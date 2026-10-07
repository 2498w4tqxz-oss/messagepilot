#!/usr/bin/env python3
import pathlib, zipfile, gzip, subprocess, json
from PIL import Image
from openpyxl import Workbook
from pptx import Presentation
p = pathlib.Path('work/file-fixtures').resolve()
for ext in ['txt', 'csv', 'rtf', 'html', 'htm', 'json', 'xml']:
    text = {'txt': 'MessagePilot file fixture', 'csv': 'name,value\nfixture,42', 'rtf': '{\\rtf1\\ansi MessagePilot fixture}', 'html': '<h1>Fixture</h1>', 'htm': '<p>Fixture</p>', 'json': '{"fixture":true}', 'xml': '<fixture>MessagePilot</fixture>'}[ext]
    (p / f'sample.{ext}').write_text(text)
for ext in ['jpg', 'jpeg', 'png', 'gif', 'webp', 'tiff', 'tif', 'bmp']:
    Image.new('RGB', (120, 80), '#1267b8').save(p / f'sample.{ext}')
Image.new('RGB', (120, 80), '#1267b8').save(p / 'sample.pdf')
(p / 'sample.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" width="120" height="80"><rect width="120" height="80" fill="blue"/></svg>')
subprocess.run(['/usr/bin/sips', '-s', 'format', 'heic', str(p / 'sample.png'), '--out', str(p / 'sample.heic')], stdout=subprocess.DEVNULL, check=True)
w = Workbook()
w.active.append(['MessagePilot', 42])
w.save(p / 'sample.xlsx')
r = Presentation()
slide = r.slides.add_slide(r.slide_layouts[0])
slide.shapes.title.text = 'MessagePilot fixture'
r.save(p / 'sample.pptx')
subprocess.run(['/opt/homebrew/bin/soffice', '-env:UserInstallation=' + str((p / 'lo').as_uri()), '--headless', '--convert-to', 'docx', '--outdir', str(p), str(p / 'sample.rtf')], check=True, stdout=subprocess.DEVNULL)
for source, ext in [('docx', 'doc'), ('xlsx', 'xls'), ('pptx', 'ppt')]:
    subprocess.run(['/opt/homebrew/bin/soffice', '-env:UserInstallation=' + str((p / 'lo').as_uri()), '--headless', '--convert-to', ext, '--outdir', str(p), str(p / f'sample.{source}')], check=True, stdout=subprocess.DEVNULL)
with zipfile.ZipFile(p / 'sample.zip', 'w') as z:
    z.write(p / 'sample.txt', 'fixture.txt')
with gzip.open(p / 'sample.gz', 'wb') as f:
    f.write(b'MessagePilot gzip fixture')
subprocess.run(['/usr/bin/tar', '-cf', str(p / 'sample.tar'), '-C', str(p), 'sample.txt'], check=True)
subprocess.run(['/usr/bin/tar', '--format', '7zip', '-cf', str(p / 'sample.7z'), '-C', str(p), 'sample.txt'], check=True)
for ext in ['mp3', 'wav', 'm4a', 'aac', 'flac']:
    subprocess.run(['/opt/homebrew/bin/ffmpeg', '-v', 'error', '-f', 'lavfi', '-i', 'sine=frequency=440:duration=0.3', '-y', str(p / f'sample.{ext}')], check=True)
for ext in ['mp4', 'mov', 'avi', 'mkv', 'webm', 'wmv']:
    subprocess.run(['/opt/homebrew/bin/ffmpeg', '-v', 'error', '-f', 'lavfi', '-i', 'color=c=blue:s=120x80:d=0.3', '-y', str(p / f'sample.{ext}')], check=True)
print('Generated synthetic formats:', len(list(p.glob('sample.*'))))
